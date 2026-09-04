"""Exportacion a Excel: el formato largo y las tablas consolidadas."""

import pandas as pd

from .agregacion import formatear_con_error

__all__ = ["exportar_largo_excel", "exportar_consolidado_excel"]


def _formatear_hojas(writer, ancho_max=60):
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter
    for hoja in writer.book.worksheets:
        for j, col in enumerate(hoja.columns, start=1):
            largo = max((len(str(c.value)) for c in col if c.value), default=10)
            hoja.column_dimensions[get_column_letter(j)].width = min(max(12, largo + 2), ancho_max)
        for fila in hoja.iter_rows(min_row=2):
            for celda in fila:
                celda.alignment = Alignment(wrap_text=True, vertical="top")


def _guardar(nombre_archivo, hojas):
    with pd.ExcelWriter(nombre_archivo, engine="openpyxl") as writer:
        for nombre, (df, con_indice) in hojas.items():
            df.to_excel(writer, sheet_name=nombre, index=con_indice)
        _formatear_hojas(writer)
    print(f"Excel guardado: {nombre_archivo}")
    return nombre_archivo


def exportar_largo_excel(largo, nombre_archivo, evidencia=None):
    """Observaciones crudas (una fila por respuesta) y, si se pasa, las citas del modo C."""
    hojas = {"Observaciones": (largo, False)}
    if evidencia is not None:
        hojas["Evidencia"] = (evidencia, False)
    return _guardar(nombre_archivo, hojas)


def exportar_consolidado_excel(cons, nombre_archivo):
    """
    Tablas de `consolidar`. Hojas: Dimensiones (media/error/n/N-A), Dimensiones ±,
    Preguntas, Preguntas ±, Observaciones.
    """
    momentos, e = list(cons["dimensiones"].columns), cons["error"]

    dim = pd.concat({"media": cons["dimensiones"], e: cons["dimensiones_err"],
                     "n": cons["n_total"], "N/A": cons["n_na"]}, axis=1).swaplevel(axis=1)
    dim = dim.reindex(columns=[(m, c) for m in momentos for c in ("media", e, "n", "N/A")])
    dim.columns = [f"{m} {c}" for m, c in dim.columns]

    filas_q, filas_txt = [], []
    for d, media in cons["preguntas"].items():
        err, n = cons["preguntas_err"][d], cons["preguntas_n"][d]
        txt = formatear_con_error(media, err)
        for q in media.index:
            filas_q.append({"Dimension": d, "Pregunta": q,
                            **{f"{m} {c}": t.loc[q, m] for m in momentos
                               for c, t in (("media", media), (e, err), ("n", n))}})
            filas_txt.append({"Dimension": d, "Pregunta": q, **{m: txt.loc[q, m] for m in momentos}})

    return _guardar(nombre_archivo, {
        "Dimensiones": (dim, True),
        "Dimensiones ±": (formatear_con_error(cons["dimensiones"], cons["dimensiones_err"]), True),
        "Preguntas": (pd.DataFrame(filas_q), False),
        "Preguntas ±": (pd.DataFrame(filas_txt), False),
        "Observaciones": (cons["largo"], False),
    })
